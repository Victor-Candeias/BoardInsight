"""Relatórios XLSX e PDF (F2), gerados no worker a partir dos filtros dos cartões.

O pedido HTTP só valida e enfileira (`report_exports` + job `report_export`); o ficheiro é
escrito em `EXPORT_DIR/<exportId>.<formato>` e descarregado depois.
"""

from collections.abc import Callable
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
from typing import Annotated, Any, Literal
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from pymongo.database import Database

from app.core.config import Settings
from app.models.cards import PRIORITY_LABELS, STATUS_LABELS
from app.repositories.base import utcnow
from app.repositories.jobs import JobRepository
from app.repositories.report_exports import ReportExportRepository
from app.services.card_query import CardEnricher, CardFilters, StateFilter, build_query, sort_spec
from app.services.dashboard import build_dashboard
from app.services.export import COLUMNS, export_row
from app.services.periods import Period, default_period

ReportFormat = Literal["xlsx", "pdf"]
StatusValue = Literal["backlog", "todo", "in_progress", "done", "archived", "discarded", "unknown"]

MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
COLUMN_BY_KEY = {c[0]: c for c in COLUMNS}
PDF_DEFAULT_COLUMNS = [
    "cardId", "title", "status", "columnName", "ownerName", "priority", "deadline", "isBlocked", "cycleTimeDays", "agingDays",
]  # fmt: skip
LIST_PROJECTION = {"_id": 0, "raw": 0, "transitions": 0}
LIST_FIELDS = ("status", "columnId", "laneId", "ownerUserId", "typeId", "priority")
KPI_LABELS = {
    "total": "Cartões (filtros atuais)",
    "backlog": "Backlog",
    "todo": "A fazer",
    "inProgress": "Em curso",
    "done": "Concluídos",
    "blocked": "Bloqueados",
    "overdue": "Em atraso",
    "dueSoon": "Prazo próximo",
    "unassigned": "Sem responsável",
}


class ReportError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


class ReportFilters(BaseModel):
    """Filtros dos cartões com os mesmos nomes dos parâmetros de query de `/cards`."""

    model_config = ConfigDict(extra="forbid")

    q: str | None = Field(None, max_length=200)
    state: StateFilter = "active"
    status: list[StatusValue] = []
    columnId: list[int] = []
    laneId: list[int] = []
    ownerUserId: list[int] = []
    unassigned: bool = False
    typeId: list[int] = []
    priority: list[Annotated[int, Field(ge=1, le=4)]] = []
    blocked: bool | None = None
    overdue: bool = False
    dueSoon: bool = False
    createdFrom: date | None = None
    createdTo: date | None = None
    endedFrom: date | None = None
    endedTo: date | None = None
    includeMissing: bool = False

    @field_validator(*LIST_FIELDS, mode="before")
    @classmethod
    def _as_list(cls, value: object) -> object:
        if value is None or value == "":
            return []
        return value if isinstance(value, list) else [value]

    @model_validator(mode="after")
    def _ranges(self) -> "ReportFilters":
        for start, end in ((self.createdFrom, self.createdTo), (self.endedFrom, self.endedTo)):
            if start and end and start > end:
                raise ValueError("intervalo de datas inválido")
        return self

    def to_card_filters(self) -> CardFilters:
        return CardFilters(
            q=self.q or None,
            state=self.state,
            status=list(self.status),
            column_id=self.columnId,
            lane_id=self.laneId,
            owner_user_id=self.ownerUserId,
            unassigned=self.unassigned,
            type_id=self.typeId,
            priority=self.priority,
            blocked=self.blocked,
            overdue=self.overdue,
            due_soon=self.dueSoon,
            created_from=self.createdFrom,
            created_to=self.createdTo,
            ended_from=self.endedFrom,
            ended_to=self.endedTo,
            include_missing=self.includeMissing,
        )


class ReportParams(BaseModel):
    """Parâmetros de um relatório (reutilizáveis em definições guardadas)."""

    model_config = ConfigDict(extra="forbid", populate_by_name=True)

    title: str | None = Field(None, max_length=120)
    filters: ReportFilters = Field(default_factory=ReportFilters)
    date_from: date | None = Field(None, alias="from")
    date_to: date | None = Field(None, alias="to")
    sort: str | None = None
    columns: list[str] | None = None

    @field_validator("sort")
    @classmethod
    def _valid_sort(cls, value: str | None) -> str | None:
        sort_spec(value)  # ValueError → erro de validação
        return value

    @field_validator("columns")
    @classmethod
    def _valid_columns(cls, value: list[str] | None) -> list[str] | None:
        if value is None:
            return None
        if not value or len(set(value)) != len(value):
            raise ValueError("colunas vazias ou repetidas")
        if unknown := [c for c in value if c not in COLUMN_BY_KEY]:
            raise ValueError(f"colunas desconhecidas: {', '.join(unknown)}")
        return value

    @model_validator(mode="after")
    def _period(self) -> "ReportParams":
        if self.date_from and self.date_to and self.date_from > self.date_to:
            raise ValueError("intervalo de datas inválido")
        return self

    def period(self, tz: ZoneInfo) -> Period:
        default_from, default_to = default_period(tz)
        return Period(self.date_from or default_from, self.date_to or default_to, tz)

    def column_specs(self, fmt: ReportFormat) -> list[tuple[str, str, str]]:
        keys = self.columns or (PDF_DEFAULT_COLUMNS if fmt == "pdf" else [c[0] for c in COLUMNS])
        return [COLUMN_BY_KEY[k] for k in keys]


def row_limit(settings: Settings, fmt: ReportFormat) -> int:
    return settings.export_pdf_max_rows if fmt == "pdf" else settings.export_max_rows


def describe_filters(f: ReportFilters, enrich: CardEnricher) -> str:
    def names(values: list[Any], label: Callable[[Any], str | None]) -> str:
        return ", ".join(label(v) or str(v) for v in values)

    def span(start: date | None, end: date | None) -> str:
        return f"{start.isoformat() if start else '…'} a {end.isoformat() if end else '…'}"

    parts = []
    if f.q:
        parts.append(f'Pesquisa: "{f.q}"')
    if f.state != "active":
        parts.append("Estado no Businessmap: " + {"archived": "arquivados", "all": "todos"}[f.state])
    if f.status:
        parts.append("Estado: " + names(f.status, STATUS_LABELS.get))
    if f.columnId:
        parts.append("Coluna: " + names(f.columnId, enrich.column_name))
    if f.laneId:
        parts.append("Lane: " + names(f.laneId, enrich.lanes.get))
    if f.ownerUserId:
        parts.append("Responsável: " + names(f.ownerUserId, enrich.user))
    if f.unassigned:
        parts.append("Sem responsável")
    if f.typeId:
        parts.append("Tipo: " + names(f.typeId, lambda t: enrich.types.get(t) or f"Tipo {t}"))
    if f.priority:
        parts.append("Prioridade: " + names(f.priority, PRIORITY_LABELS.get))
    if f.blocked is not None:
        parts.append("Bloqueados" if f.blocked else "Não bloqueados")
    if f.overdue:
        parts.append("Em atraso")
    if f.dueSoon:
        parts.append("Prazo próximo")
    if f.createdFrom or f.createdTo:
        parts.append("Criados: " + span(f.createdFrom, f.createdTo))
    if f.endedFrom or f.endedTo:
        parts.append("Concluídos: " + span(f.endedFrom, f.endedTo))
    if f.includeMissing:
        parts.append("Inclui cartões em falta")
    return "; ".join(parts) or "Sem filtros"


def _local(dt: datetime | None, tz: ZoneInfo) -> str:
    return dt.astimezone(tz).strftime("%Y-%m-%d %H:%M") if dt else "—"


def build_report_data(
    db: Database, settings: Settings, board: dict[str, Any], params: ReportParams, fmt: ReportFormat, now: datetime
) -> dict[str, Any]:
    tz = settings.tz
    filters = params.filters.to_card_filters()
    period = params.period(tz)
    query = build_query(board["boardId"], filters, tz)
    total = db.cards.count_documents(query)
    limit = row_limit(settings, fmt)
    if total > limit:
        raise ReportError("export_too_large", f"{total} cartões excedem o máximo de {limit} para {fmt.upper()}; refine os filtros")
    enrich = CardEnricher(db, board, now)
    rows = [export_row(enrich(c), now) for c in db.cards.find(query, LIST_PROJECTION).sort(sort_spec(params.sort))]
    dashboard = build_dashboard(db, board, filters, period, "week", now)
    title = params.title or "Relatório geral"
    meta = [
        ("Relatório", title),
        ("Board", f"{board.get('name') or 'Board'} ({board['boardId']})"),
        ("Período", f"{period.start.isoformat()} a {period.end.isoformat()}"),
        ("Filtros", describe_filters(params.filters, enrich)),
        ("Gerado em", _local(now, tz)),
        ("Dados de", _local(board.get("lastSuccessfulSyncAt"), tz) if board.get("lastSuccessfulSyncAt") else "sem sincronização"),
    ]
    kpis = [(label, dashboard["kpis"].get(key, 0)) for key, label in KPI_LABELS.items()]
    kpis += [
        ("Concluídos no período", sum(t["completed"] for t in dashboard["timeline"])),
        ("Criados no período", sum(t["created"] for t in dashboard["timeline"])),
    ]
    return {"title": title, "meta": meta, "kpis": kpis, "rows": rows, "columns": params.column_specs(fmt)}


# --- XLSX -------------------------------------------------------------------------------------


def _xlsx_cell(ws: Any, value: Any, kind: str, tz: ZoneInfo) -> Any:
    from openpyxl.cell import WriteOnlyCell

    if value is None:
        return None
    if kind == "bool":
        return "Sim" if value else "Não"
    if kind == "datetime":
        cell = WriteOnlyCell(ws, value.astimezone(tz).replace(tzinfo=None))
        cell.number_format = "yyyy-mm-dd hh:mm"
        return cell
    if kind in ("int", "number"):
        return value
    cell = WriteOnlyCell(ws, ", ".join(value) if kind == "list" else str(value))
    cell.data_type = "s"  # nunca interpretar como fórmula
    return cell


def write_xlsx(path: Path, data: dict[str, Any], tz: ZoneInfo) -> None:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)

    def header(ws: Any, text: str) -> Any:
        cell = WriteOnlyCell(ws, text)
        cell.font = bold
        return cell

    wb = Workbook(write_only=True)
    summary = wb.create_sheet("Resumo")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 90
    for label, value in data["meta"]:
        summary.append([header(summary, label), _xlsx_cell(summary, value, "text", tz)])
    summary.append([])
    summary.append([header(summary, "Indicador"), header(summary, "Valor")])
    for label, value in data["kpis"]:
        summary.append([label, value])

    columns, rows = data["columns"], data["rows"]
    sheet = wb.create_sheet("Cartões")
    sheet.freeze_panes = "A2"
    for i, (key, _, kind) in enumerate(columns, start=1):
        width = 50 if key == "title" else 18 if kind in ("datetime", "text", "list") else 12
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    sheet.append([header(sheet, h) for _, h, _ in columns])
    for row in rows:
        sheet.append([_xlsx_cell(sheet, row[k], kind, tz) for k, _, kind in columns])
    wb.save(path)


# --- PDF --------------------------------------------------------------------------------------

FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
]
COLUMN_WEIGHTS = {"title": 5.0, "blockReason": 3.0}
KIND_WEIGHTS = {"text": 2.0, "list": 2.4, "datetime": 2.2, "number": 1.2, "int": 1.1, "bool": 1.4}


@lru_cache
def _fonts() -> tuple[str, str]:
    """DejaVu (Unicode completo) se existir; senão Helvetica (Latin-1, cobre o português)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in FONT_CANDIDATES:
        if Path(regular).is_file() and Path(bold).is_file():
            pdfmetrics.registerFont(TTFont("BI-Regular", regular))
            pdfmetrics.registerFont(TTFont("BI-Bold", bold))
            return "BI-Regular", "BI-Bold"
    return "Helvetica", "Helvetica-Bold"


def _pdf_text(value: Any, kind: str, tz: ZoneInfo) -> str:
    if value is None:
        return ""
    if kind == "bool":
        return "Sim" if value else "Não"
    if kind == "datetime":
        return value.astimezone(tz).strftime("%Y-%m-%d %H:%M")
    if kind == "number":
        return str(value).replace(".", ",")
    if kind == "list":
        return ", ".join(value)
    return str(value)


def write_pdf(path: Path, data: dict[str, Any], tz: ZoneInfo) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _fonts()
    page_w, page_h = landscape(A4)
    margin = 12 * mm
    width = page_w - 2 * margin
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=11, leading=14, spaceBefore=8, spaceAfter=4)
    body = ParagraphStyle("body", fontName=font, fontSize=8, leading=10)
    cell = ParagraphStyle("cell", fontName=font, fontSize=7, leading=8.5)
    head = ParagraphStyle("head", fontName=bold, fontSize=7, leading=8.5)
    grid = [
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]

    def p(text: str, style: ParagraphStyle = cell) -> Paragraph:
        return Paragraph(escape(text), style)

    story: list[Any] = [Paragraph(escape(data["title"]), h1)]
    meta = Table([[p(k, head), p(str(v), body)] for k, v in data["meta"][1:]], colWidths=[35 * mm, width - 35 * mm])
    meta.setStyle(TableStyle(grid))
    story += [meta, Paragraph("Indicadores", h2)]
    kpis = data["kpis"]
    kpi_table = Table([[p(k, head) for k, _ in kpis], [p(str(v), body) for _, v in kpis]], colWidths=[width / len(kpis)] * len(kpis))
    kpi_table.setStyle(TableStyle([*grid, ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9"))]))
    story.append(kpi_table)

    columns, rows = data["columns"], data["rows"]
    story.append(Paragraph(f"Cartões ({len(rows)})", h2))
    if rows:
        weights = [COLUMN_WEIGHTS.get(key, KIND_WEIGHTS.get(kind, 1.5)) for key, _, kind in columns]
        col_widths = [width * w / sum(weights) for w in weights]
        table_rows = [[p(h, head) for _, h, _ in columns]]
        table_rows += [[p(_pdf_text(row[k], kind, tz)) for k, _, kind in columns] for row in rows]
        table = LongTable(table_rows, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    *grid,
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("Nenhum cartão cumpre os filtros.", body))
    story.append(Spacer(1, 4 * mm))

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(margin, 7 * mm, f"BoardInsight · {data['title']} · uso interno")
        canvas.drawRightString(page_w - margin, 7 * mm, f"Página {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=(page_w, page_h),
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=14 * mm,
        title=data["title"],
        author="BoardInsight",
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


WRITERS: dict[str, Callable[[Path, dict[str, Any], ZoneInfo], None]] = {"xlsx": write_xlsx, "pdf": write_pdf}


def enqueue_export(
    db: Database,
    settings: Settings,
    board_id: int,
    fmt: ReportFormat,
    params: ReportParams,
    *,
    definition_id: str | None = None,
    default_title: str = "Relatório geral",
) -> dict[str, Any]:
    """Verifica o limite de linhas e cria `report_exports` + job `report_export`."""
    total = db.cards.count_documents(build_query(board_id, params.filters.to_card_filters(), settings.tz))
    limit = row_limit(settings, fmt)
    if total > limit:
        raise ReportError("export_too_large", f"{total} cartões excedem o máximo de {limit} para {fmt.upper()}; refine os filtros")
    title = params.title or default_title
    stored = params.model_dump(mode="json", by_alias=True)
    stored["title"] = title
    repo = ReportExportRepository(db)
    export = repo.create(
        {"boardId": board_id, "format": fmt, "report": "general", "title": title, "definitionId": definition_id, "params": stored},
        timedelta(days=settings.export_retention_days),
    )
    job, _ = JobRepository(db).enqueue("report_export", {"exportId": export["id"]}, dedupe=False)
    repo.update(export["id"], {"jobId": job["id"]})
    return {**export, "jobId": job["id"]}


def export_path(settings: Settings, export: dict[str, Any]) -> Path:
    return settings.export_dir / f"{export['id']}.{export['format']}"


def run_export(db: Database, settings: Settings, export_id: str) -> dict[str, Any]:
    """Gera o ficheiro de um `report_exports` (executado pelo worker)."""
    repo = ReportExportRepository(db)
    export = repo.get(export_id)
    if export is None:
        raise ReportError("export_not_found", "exportação inexistente ou expirada")
    repo.update(export_id, {"status": "running", "startedAt": utcnow(), "error": None})
    try:
        board = db.boards.find_one({"boardId": export["boardId"]})
        if board is None:
            raise ReportError("board_not_found", f"o board {export['boardId']} não está configurado")
        now = utcnow()
        fmt = export["format"]
        data = build_report_data(db, settings, board, ReportParams.model_validate(export["params"]), fmt, now)
        settings.export_dir.mkdir(parents=True, exist_ok=True)
        path = export_path(settings, export)
        tmp = path.with_name(path.name + ".tmp")
        WRITERS[fmt](tmp, data, settings.tz)
        tmp.replace(path)
    except ReportError as exc:
        repo.update(export_id, {"status": "error", "finishedAt": utcnow(), "error": {"code": exc.code, "message": exc.message}})
        raise
    except Exception as exc:
        repo.update(export_id, {"status": "error", "finishedAt": utcnow(), "error": {"code": "internal_error", "message": type(exc).__name__}})
        raise
    stamp = now.astimezone(settings.tz).strftime("%Y%m%d-%H%M")
    result = {
        "fileName": f"boardinsight-board{export['boardId']}-relatorio-{stamp}.{fmt}",
        "size": path.stat().st_size,
        "rows": len(data["rows"]),
    }
    repo.update(export_id, {"status": "done", "finishedAt": utcnow(), **result})
    return {"exportId": export_id, **result}
