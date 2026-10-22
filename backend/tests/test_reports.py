import asyncio
import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.jobs.handlers import HANDLERS
from app.jobs.worker import Worker
from app.repositories.base import utcnow

URL = "/api/v1/reports/export"


def run_worker(db, settings):
    # Só o handler de relatórios: o job de sync pendente do fixture não deve correr.
    worker = Worker(db, settings, handlers={"report_export": HANDLERS["report_export"]})
    while asyncio.run(worker.run_one()):
        pass


def test_xlsx_export_runs_in_worker_and_downloads(synced, db, settings):
    db.cards.update_one({"cardId": 5001}, {"$set": {"title": "=HYPERLINK(\"http://x\")"}})
    resp = synced.post(URL, json={"boardId": 96, "format": "xlsx", "sort": "cardId", "filters": {"status": ["backlog", "done"]}})
    assert resp.status_code == 202
    export = resp.json()
    assert export["status"] == "pending" and export["jobId"] and export["downloadUrl"] is None
    assert export["params"]["filters"]["status"] == ["backlog", "done"]
    download = f"/api/v1/reports/exports/{export['id']}/download"
    assert synced.get(download).json()["error"]["code"] == "export_not_ready"

    run_worker(db, settings)
    done = synced.get(f"/api/v1/reports/exports/{export['id']}").json()
    assert done["status"] == "done" and done["rows"] == 3 and done["downloadUrl"] == download
    assert synced.get("/api/v1/reports/exports", params={"boardId": 96}).json()[0]["id"] == export["id"]

    file = synced.get(download)
    assert file.status_code == 200 and file.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="boardinsight-board96-relatorio-' in file.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(file.content))
    assert wb.sheetnames == ["Resumo", "Cartões"]
    summary = {r[0]: r[1] for r in wb["Resumo"].iter_rows(values_only=True) if r and r[0]}
    assert summary["Filtros"] == "Estado: Backlog, Concluído" and summary["Cartões (filtros atuais)"] == 3
    rows = list(wb["Cartões"].iter_rows(values_only=True))
    header = rows[0]
    assert header[:3] == ("ID", "ID personalizado", "Título") and [r[0] for r in rows[1:]] == [5001, 5005, 5006]
    title_cell = wb["Cartões"].cell(row=2, column=3)
    assert title_cell.value.startswith("=HYPERLINK") and title_cell.data_type == "s"
    ended = rows[3][header.index("Fim")]
    assert ended == datetime(2026, 9, 30, 13, 0)  # hora local de Lisboa
    assert rows[3][header.index("Cycle time (dias)")] == 8.12


def test_pdf_export_with_selected_columns(synced, db, settings):
    body = {"boardId": 96, "format": "pdf", "title": "Ana — semana", "filters": {"ownerUserId": "11"}, "columns": ["cardId", "title", "deadline"]}
    export = synced.post(URL, json=body).json()
    run_worker(db, settings)
    done = synced.get(f"/api/v1/reports/exports/{export['id']}").json()
    assert done["status"] == "done" and done["rows"] == 2 and done["title"] == "Ana — semana"
    file = synced.get(done["downloadUrl"])
    assert file.headers["content-type"] == "application/pdf" and file.content.startswith(b"%PDF")
    assert (settings.export_dir / f"{export['id']}.pdf").stat().st_size == done["size"]


def test_export_validation_and_limits(synced, client, settings):
    assert synced.post(URL, json={"boardId": 1, "format": "pdf"}).json()["error"]["code"] == "board_not_found"
    for bad in (
        {"boardId": 96, "format": "docx"},
        {"boardId": 96, "format": "pdf", "columns": ["nope"]},
        {"boardId": 96, "format": "pdf", "filters": {"status": "weird"}},
        {"boardId": 96, "format": "pdf", "sort": "nope"},
        {"boardId": 96, "format": "pdf", "from": "2026-10-10", "to": "2026-10-01"},
    ):
        assert synced.post(URL, json=bad).status_code == 422, bad
    settings.export_pdf_max_rows = 3
    resp = synced.post(URL, json={"boardId": 96, "format": "pdf"})
    assert resp.status_code == 422 and resp.json()["error"]["code"] == "export_too_large"
    assert synced.post(URL, json={"boardId": 96, "format": "xlsx"}).status_code == 202


def test_failed_expired_and_missing_exports(synced, db, settings):
    assert synced.get("/api/v1/reports/exports/nope").status_code == 404
    export = synced.post(URL, json={"boardId": 96, "format": "xlsx"}).json()
    # Worker morreu: o job passou a erro sem atualizar a exportação.
    from bson import ObjectId

    db.jobs.update_one({"_id": ObjectId(export["jobId"])}, {"$set": {"status": "error", "error": {"code": "timeout", "message": "x"}}})
    failed = synced.get(f"/api/v1/reports/exports/{export['id']}").json()
    assert failed["status"] == "error" and failed["error"]["code"] == "timeout"
    assert synced.get(f"/api/v1/reports/exports/{export['id']}/download").json()["error"]["code"] == "export_failed"

    other = synced.post(URL, json={"boardId": 96, "format": "xlsx"}).json()
    run_worker(db, settings)
    (settings.export_dir / f"{other['id']}.xlsx").unlink()
    assert synced.get(f"/api/v1/reports/exports/{other['id']}/download").status_code == 410
    db.report_exports.update_one({"_id": ObjectId(other["id"])}, {"$set": {"expiresAt": utcnow() - timedelta(seconds=1)}})
    # O índice TTL remove o documento (o 410 cobre o atraso do TTL do MongoDB).
    assert synced.get(f"/api/v1/reports/exports/{other['id']}/download").json()["error"]["code"] == "export_not_found"
